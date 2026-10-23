import re, sys
from docx import Document
from docx.shared import Pt, Inches
base="/home/hatch/workspace/publications/weekly-series/meta-analysis-llm-citation-fabrication/review2/"
doc=Document()
for s in doc.sections: s.left_margin=Inches(0.9); s.right_margin=Inches(0.9)
lines=open(base+"manuscript.md",encoding="utf-8").read().splitlines()
i=0
def clean(t): return re.sub(r'\*\*(.*?)\*\*',r'\1',t).strip()
while i<len(lines):
    line=lines[i]
    if line.startswith("|") and i+1<len(lines) and lines[i+1].startswith("|---"):
        rows=[]
        while i<len(lines) and lines[i].startswith("|"):
            rows.append([clean(c) for c in lines[i].strip().strip('|').split('|')]); i+=1
        rows=[r for r in rows if not all(set(c)<=set('-: ') for c in r)]
        t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'
        for ri,r in enumerate(rows):
            for ci,c in enumerate(r): t.cell(ri,ci).text=c
        continue
    if line.startswith("# "): doc.add_heading(clean(line[2:]),level=0)
    elif line.startswith("## "): doc.add_heading(clean(line[3:]),level=1)
    elif line.startswith("### "): doc.add_heading(clean(line[4:]),level=2)
    elif line.strip():
        p=doc.add_paragraph()
        # bold segments
        parts=re.split(r'(\*\*.*?\*\*)',line)
        for part in parts:
            run=p.add_run(part[2:-2] if part.startswith("**") else part)
            if part.startswith("**"): run.bold=True
            run.font.size=Pt(11)
    i+=1
doc.add_page_break(); doc.add_heading("Forest plot: completely fabricated",level=1)
doc.add_picture(base+"forest-fabricated.png",width=Inches(6.5))
doc.add_heading("Forest plot: fully accurate",level=1)
doc.add_picture(base+"forest-fully_accurate.png",width=Inches(6.5))
doc.add_heading("PRISMA flow diagram",level=1)
doc.add_picture(base+"prisma-flow-diagram.png",width=Inches(5.5))
doc.save(base+"manuscript.docx")
print("docx built")
